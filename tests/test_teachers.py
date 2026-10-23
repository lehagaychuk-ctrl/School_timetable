import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / '.vendor')]

import json
import sqlite3
import tempfile
import unittest
from copy import deepcopy
from io import BytesIO

from app import create_app, empty_day
from openpyxl import load_workbook
from docx import Document
from pypdf import PdfReader


class TeacherCatalogTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.path = Path(self.tmp.name) / 'teachers.sqlite3'
        self.client = create_app(self.path).test_client()
        self.day_date = '2026-10-15'

    def tearDown(self):
        self.tmp.cleanup()

    def add(self, name='Иванова А.А.', subjects=None, rooms=None):
        result = self.client.post('/api/teachers', json={'name': name, 'subjects': subjects or ['Математика'], 'rooms': rooms or ['14']})
        self.assertEqual(result.status_code, 201, result.json)
        return result.json['teacher']

    def lesson(self, teacher, cid=1, slot=1):
        return {'class_id': cid, 'slot': slot, 'groups': [{'teacher_id': teacher['id'], 'teacher': 'Клиентское имя игнорируется',
                 'subject': 'Математика', 'rooms': [str(100 + cid)], 'label': ''}], 'cancelled': False, 'note': ''}

    def save(self, lessons):
        body = empty_day()
        body['lessons'] = lessons
        revision = self.client.get('/api/day/' + self.day_date).json['revision']
        result = self.client.put('/api/day/' + self.day_date, json={'day': body, 'revision': revision})
        self.assertEqual(result.status_code, 200, result.json)
        return result.json

    def test_catalog_validation_duplicate_and_persistence(self):
        teacher = self.add(subjects='Математика\nАлгебра\nматематика', rooms='14,25')
        self.assertEqual(teacher['subjects'], ['Математика', 'Алгебра'])
        self.assertEqual(teacher['rooms'], ['14', '25'])
        self.assertEqual(self.client.post('/api/teachers', json={'name': '  иванова   а.а.  '}).status_code, 409)
        for body in [{'name': ''}, {'name': 'Учитель', 'rooms': ['1', '2', '3', '4', '5']},
                     {'name': 'Учитель', 'subjects': 123}]:
            self.assertEqual(self.client.post('/api/teachers', json=body).status_code, 400)
        restarted = create_app(self.path).test_client().get('/api/state').json
        self.assertEqual(restarted['teachers'], [teacher])

    def test_stable_ids_conflicts_and_unknown_reference(self):
        teacher = self.add()
        result = self.save([self.lesson(teacher, 1, 8), self.lesson(teacher, 25, 1)])
        self.assertTrue(any('Учитель' in i['message'] for i in result['issues']))
        self.assertEqual(result['day']['lessons'][0]['groups'][0]['teacher'], teacher['name'])
        bad = empty_day()
        bad['lessons'] = [self.lesson({**teacher, 'id': 9999})]
        self.assertEqual(self.client.put('/api/day/' + self.day_date, json={'day': bad, 'revision': result['revision']}).status_code, 400)

    def test_rename_updates_dates_weekly_and_all_exports(self):
        teacher = self.add()
        saved = self.save([self.lesson(teacher)])
        self.client.post(f'/api/day/{self.day_date}/weekly', json={'revision': saved['revision']})
        self.client.post(f'/api/day/{self.day_date}/approve', json={'revision': saved['revision']})
        result = self.client.put(f'/api/teachers/{teacher["id"]}', json={**teacher, 'name': 'Переименованная Б.Б.', 'rooms': ['99']})
        self.assertEqual(result.status_code, 200)
        current = self.client.get('/api/day/' + self.day_date).json
        group = current['day']['lessons'][0]['groups'][0]
        self.assertEqual(group['teacher_id'], teacher['id'])
        self.assertEqual(group['teacher'], 'Переименованная Б.Б.')
        self.assertEqual(group['rooms'], ['101'])
        self.assertFalse(current['day']['approved'])
        self.assertGreater(current['revision'], saved['revision'])
        self.client.post('/api/day/2026-10-22/copy', json={'source': 'weekly', 'revision': 0})
        copied = self.client.get('/api/day/2026-10-22').json['day']['lessons'][0]['groups'][0]
        self.assertEqual(copied['teacher'], 'Переименованная Б.Б.')
        settings = self.client.get('/api/state').json['settings']
        settings['template']['show_teachers'] = True
        self.client.put('/api/settings', json=settings)
        for fmt in ('pdf', 'xlsx', 'docx'):
            result = self.client.get(f'/api/export/{self.day_date}/{fmt}?scope=class&class_id=1')
            self.assertEqual(result.status_code, 200)
            if fmt == 'pdf':
                text = ' '.join(p.extract_text() for p in PdfReader(BytesIO(result.data)).pages)
            elif fmt == 'xlsx':
                book = load_workbook(BytesIO(result.data))
                text = ' '.join(str(c.value) for ws in book for row in ws for c in row if c.value)
            else:
                doc = Document(BytesIO(result.data))
                text = ' '.join(c.text for table in doc.tables for row in table.rows for c in row.cells)
            self.assertIn('Переименованная Б.Б.', ' '.join(text.split()))

    def test_archive_restore_and_stale_edits(self):
        teacher = self.add()
        self.save([self.lesson(teacher)])
        result = self.client.put(f'/api/teachers/{teacher["id"]}', json={**teacher, 'active': False})
        archived = result.json['teacher']
        self.assertFalse(archived['active'])
        current = self.client.get('/api/day/' + self.day_date).json
        self.assertEqual(len(current['day']['lessons']), 1)
        self.assertTrue(any('в архиве' in i['message'] for i in current['issues']))
        self.assertEqual(self.client.put(f'/api/teachers/{teacher["id"]}', json=teacher).status_code, 409)
        self.assertEqual(self.client.post('/api/teachers', json={'name': teacher['name']}).status_code, 409)
        restored = self.client.put(f'/api/teachers/{teacher["id"]}', json={**archived, 'active': True}).json['teacher']
        self.assertEqual(restored['id'], teacher['id'])
        self.assertFalse(self.client.get('/api/day/' + self.day_date).json['issues'])

    def test_existing_free_text_becomes_linked_when_added_to_catalog(self):
        body = empty_day()
        body['lessons'] = [{'class_id': 1, 'slot': 1, 'groups': [{'teacher': ' Иванова  А.А. ', 'subject': 'Русский язык', 'rooms': ['7']}]}]
        self.client.put('/api/day/' + self.day_date, json={'day': body, 'revision': 0})
        teacher = self.add()
        group = self.client.get('/api/day/' + self.day_date).json['day']['lessons'][0]['groups'][0]
        self.assertEqual(group['teacher_id'], teacher['id'])
        self.assertEqual(group['rooms'], ['7'])
        self.assertEqual(group['subject'], 'Русский язык')

    def test_migration_of_legacy_database_is_idempotent(self):
        legacy = Path(self.tmp.name) / 'legacy.sqlite3'
        body = empty_day()
        body['lessons'] = [{'class_id': 1, 'slot': 1, 'groups': [{'teacher': 'Иванова А.А.', 'subject': 'Русский язык', 'rooms': ['7']}]},
                           {'class_id': 2, 'slot': 2, 'groups': [{'teacher': ' иванова   а.а. ', 'subject': 'Литература', 'rooms': ['8']}]}]
        with sqlite3.connect(legacy) as conn:
            conn.executescript('CREATE TABLE days(date TEXT PRIMARY KEY, body TEXT NOT NULL, revision INTEGER NOT NULL); CREATE TABLE weekly(weekday INTEGER PRIMARY KEY, body TEXT NOT NULL);')
            conn.execute('INSERT INTO days VALUES (?,?,?)', (self.day_date, json.dumps(body), 3))
            conn.execute('INSERT INTO weekly VALUES (?,?)', (3, json.dumps(body)))
        conn.close()
        client = create_app(legacy).test_client()
        catalog = client.get('/api/state').json['teachers']
        self.assertEqual(len(catalog), 1)
        self.assertEqual(catalog[0]['subjects'], ['Русский язык', 'Литература'])
        self.assertEqual(catalog[0]['class_ids'], [1, 2])
        day = client.get('/api/day/' + self.day_date).json
        self.assertEqual(day['revision'], 4)
        self.assertEqual({l['groups'][0]['teacher_id'] for l in day['day']['lessons']}, {catalog[0]['id']})
        again = create_app(legacy).test_client()
        self.assertEqual(again.get('/api/day/' + self.day_date).json, day)
        self.assertEqual(again.get('/api/state').json['teachers'], catalog)


if __name__ == '__main__':
    unittest.main()
