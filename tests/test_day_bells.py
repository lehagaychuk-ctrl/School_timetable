"""Date-specific bells keep imported lesson zero accurate across all app views."""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / '.vendor')]

import tempfile
import unittest
from copy import deepcopy
from io import BytesIO
from app import create_app, empty_day
from scheduling import teaching_entries
from pypdf import PdfReader
from openpyxl import load_workbook
from docx import Document


class DayBellsTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.client = create_app(Path(self.tmp.name) / 'school.sqlite3').test_client()
        self.state = self.client.get('/api/state').json
        self.date = '2026-10-16'
        self.day = empty_day()
        self.day['bells'] = {'1': {'1': '08:00–08:40', '2': '13:15–13:55'},
                             '2': {'0': '13:15–13:55', '1': '14:00–14:40'}}
        self.day['lessons'] = [{'class_id': 25, 'slot': 0, 'groups': [
            {'subject': 'Физика', 'teacher': '', 'rooms': ['8']}]}]

    def tearDown(self):
        self.tmp.cleanup()

    def save(self, body=None):
        return self.client.put('/api/day/' + self.date, json={'day': body or self.day, 'revision': 0})

    def test_overrides_survive_save_and_copy_without_changing_other_dates(self):
        saved = self.save()
        self.assertEqual(saved.status_code, 200, saved.json)
        self.assertEqual(saved.json['day']['bells'], self.day['bells'])
        self.assertFalse(any(i['severity'] == 'error' for i in saved.json['issues']))
        self.assertEqual(self.client.get('/api/state').json['settings']['bells'], self.state['settings']['bells'])
        self.assertNotIn('bells', self.client.get('/api/day/2026-10-15').json['day'])
        copied = self.client.post('/api/day/2026-10-17/copy', json={'source': 'date', 'source_date': self.date, 'revision': 0})
        self.assertEqual(copied.status_code, 200, copied.json)
        self.assertEqual(self.client.get('/api/day/2026-10-17').json['day']['bells'], self.day['bells'])

    def test_invalid_bells_and_zero_without_a_time_are_rejected(self):
        for bells in ({}, {'2': {'0': '15:00–14:00'}}, {'2': {'0': '13:15–13:55', '1': '13:30–14:10'}},
                      {'2': {'13': '13:15–13:55'}}, {'3': {'0': '13:15–13:55'}}, {'2': []}):
            body = deepcopy(self.day)
            body['bells'] = bells
            result = self.save(body)
            self.assertEqual(result.status_code, 400, result.json)

    def test_constructor_uses_zero_time_for_availability_busy_and_weekly_workload(self):
        teacher = self.client.post('/api/teachers', json={
            'name': 'Учитель А.А.', 'subjects': ['Физика'], 'rooms': ['8'], 'class_ids': [1, 25, 26],
            'availability_enabled': True, 'availability': {'4': ['13:15–13:55']}, 'weekly_hours': 3}).json['teacher']
        self.day['lessons'][0]['groups'][0]['teacher_id'] = teacher['id']
        self.assertEqual(self.save().status_code, 200)
        for cid, slot in [(1, 2), (26, 0)]:
            result = self.client.get(f'/api/constructor/{self.date}?class_id={cid}&slot={slot}')
            self.assertEqual(result.status_code, 200, result.json)
            self.assertEqual(result.json['time'], '13:15–13:55')
            choice = result.json['candidates'][0]
            self.assertEqual(choice['week_count'], 1)
            self.assertFalse(choice['available'])
            self.assertEqual(choice['busy_lessons'][0]['slot'], 0)
            self.assertEqual(choice['busy_lessons'][0]['time'], '13:15–13:55')
        outside = self.client.get(f'/api/constructor/{self.date}?class_id=26&slot=1').json['candidates'][0]
        self.assertFalse(outside['available'])
        self.assertFalse(outside['busy_lessons'])
        self.assertTrue(any('часы работы' in r for r in outside['reasons']))
        self.assertEqual(self.client.get('/api/constructor/2026-10-15?class_id=26&slot=0').status_code, 400)

    def test_all_exports_keep_number_zero_and_date_times(self):
        self.assertEqual(self.save().status_code, 200)
        entries = list(teaching_entries(self.day, self.state['classes'], self.state['settings']['bells']))
        self.assertEqual(entries[0]['interval'], (795, 835))
        for fmt in ('pdf', 'xlsx', 'docx'):
            result = self.client.get(f'/api/export/{self.date}/{fmt}?scope=2&draft=1')
            self.assertEqual(result.status_code, 200)
            if fmt == 'pdf':
                text = '\n'.join(p.extract_text() for p in PdfReader(BytesIO(result.data)).pages)
            elif fmt == 'xlsx':
                book = load_workbook(BytesIO(result.data))
                text = '\n'.join(str(c.value) for sheet in book for row in sheet for c in row if c.value is not None)
            else:
                doc = Document(BytesIO(result.data))
                text = '\n'.join(c.text for table in doc.tables for row in table.rows for c in row.cells)
            self.assertIn('13:15–13:55', text)
            self.assertIn('Физика', text)
            self.assertIn('0', [line.strip() for line in text.splitlines()])
            self.assertNotIn('19:45–20:25', text)

    def test_changing_class_shift_keeps_saved_zero_visible_in_draft(self):
        self.assertEqual(self.save().status_code, 200)
        classes = deepcopy(self.state['classes'])
        next(c for c in classes if c['id'] == 25)['shift'] = 1
        self.assertEqual(self.client.put('/api/classes', json={'classes': classes}).status_code, 200)
        issues = self.client.get('/api/day/' + self.date).json['issues']
        self.assertTrue(any('нет времени' in i['message'] for i in issues))
        result = self.client.get(f'/api/export/{self.date}/pdf?scope=1&draft=1')
        self.assertEqual(result.status_code, 200)
        text = '\n'.join(p.extract_text() for p in PdfReader(BytesIO(result.data)).pages)
        self.assertIn('Физика', text)
        self.assertIn('Время не задано', ' '.join(text.split()))


if __name__ == '__main__':
    unittest.main()
