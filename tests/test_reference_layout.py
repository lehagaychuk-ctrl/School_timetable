"""Measured print layout matching the supplied school timetable photographs."""
import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
sys.path[:0] = [str(ROOT), str(ROOT / '.vendor')]

import tempfile
import unittest
from copy import deepcopy
from io import BytesIO
from zipfile import ZipFile
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from openpyxl import load_workbook
from pypdf import PdfReader
from app import create_app, empty_day, REFERENCE_TEMPLATE, default_bells
from documents import make_layout, build_document, table_blocks


def reference_fixture():
    """Independent test data only; no school data is read or changed."""
    groups = [
        (1, '5-е и инженерный', ['5' + letter for letter in 'АБВГДЕЖЗ'] + ['7А (инж)']),
        (1, '8-е и естественный', ['7Е (ест)'] + ['8' + letter for letter in 'АБВГДЕЖЗ']),
        (1, '9-е', ['9' + letter for letter in 'АБВГДЕЖЗИ']),
        (1, 'Старшие', ['10А', '10Б', '10В', '11А', '11Б', '11В']),
        (2, '6-е', ['6' + letter for letter in 'АБВГДЕЖЗ']),
        (2, '7-е', ['7' + letter for letter in 'АБВГДЖ']),
    ]
    classes = []
    for shift, group, names in groups:
        for name in names:
            cid = len(classes) + 1
            classes.append({'id': cid, 'name': name, 'shift': shift, 'position': cid, 'print_group': group})
    settings = {'template': {**REFERENCE_TEMPLATE, 'school_name': 'Тестовая школа', 'director': 'Директор И.И.'},
                'bells': {'1': ['08:00–08:40', '08:45–09:25', '09:45–10:25', '10:45–11:25', '11:45–12:25', '12:30–13:10', '13:15–13:55'],
                          '2': ['14:00–14:40', '15:00–15:40', '15:45–16:25', '16:30–17:10', '17:15–17:55', '18:00–18:40']}}
    day = empty_day()
    subjects = ['Рус.язык', 'Литература', 'Алгебра', 'История', 'География', 'Физ.культ.']
    for cls in classes:
        for slot in range(1, len(settings['bells'][str(cls['shift'])]) + 1):
            day['lessons'].append({'class_id': cls['id'], 'slot': slot,
                                   'groups': [{'subject': subjects[(slot + cls['id']) % len(subjects)],
                                               'teacher': '', 'rooms': [str(100 + cls['id'])], 'label': ''}]})
    day['print_notes'] = [{'class_id': 1, 'text': 'Английский язык — занятия всем классом'},
                          {'class_id': 10, 'text': 'Полезные навыки — 8Б, 8З'},
                          {'class_id': 28, 'text': 'Информатика — весь класс'},
                          {'class_id': 34, 'text': 'Технология — занятия всем классом'},
                          {'class_id': 42, 'text': 'Английский язык — весь класс'}]
    return classes, day, settings


class ReferenceLayoutTests(unittest.TestCase):
    def test_reference_shape_fits_four_blocks_then_two_on_two_landscape_pages(self):
        classes, day, settings = reference_fixture()
        pages = make_layout('2026-10-16', day, classes, settings, False, [])
        self.assertEqual([p['shift'] for p in pages], [1, 2])
        self.assertEqual([[len(i['rows'][0]) - 2 for i in p['items'] if i['type'] == 'table'] for p in pages], [[9, 9, 9, 6], [8, 6]])
        for page in pages:
            self.assertGreater(page['width'], page['height'])
            self.assertGreaterEqual(page['remaining'], 0)
            for table in [i for i in page['items'] if i['type'] == 'table']:
                self.assertEqual(table['rows'][0][:2], ['', ''])
                self.assertTrue(table['compact'])
                self.assertLess(max(table['heights']), 16)
        self.assertTrue(any('Утверждаю' in i['text'] and i['width'] <= 245 for i in pages[0]['header']))
        self.assertFalse(any('Утверждаю' in i['text'] for i in pages[1]['header']))
        first_table = next(i for i in pages[0]['items'] if i['type'] == 'table')
        last_table = [i for i in pages[0]['items'] if i['type'] == 'table'][-1]
        self.assertLess(sum(last_table['widths']), sum(first_table['widths']))
        items = pages[0]['items']
        self.assertEqual(items[0]['text'], day['print_notes'][0]['text'])
        self.assertTrue(items[0]['bold'])
        self.assertEqual(items[1]['rows'][0][2], '5А')

    def test_editable_exports_keep_gray_headers_left_aligned_lessons_and_complete_text(self):
        classes, day, settings = reference_fixture()
        pdf = PdfReader(BytesIO(build_document('pdf', '2026-10-16', day, classes, settings)))
        self.assertEqual(len(pdf.pages), 2)
        text = '\n'.join(p.extract_text() for p in pdf.pages)
        self.assertTrue(all(c['name'] in text for c in classes))
        self.assertTrue(all(note['text'] in text for note in day['print_notes']))
        self.assertNotIn('Страница', text)
        self.assertFalse(any(p.images for p in pdf.pages))
        document_bytes = build_document('docx', '2026-10-16', day, classes, settings)
        doc = Document(BytesIO(document_bytes))
        self.assertEqual(len(doc.tables), 6)
        self.assertEqual([len(table.columns) for table in doc.tables], [11, 11, 11, 8, 10, 8])
        self.assertEqual(doc.tables[0].cell(1, 2).paragraphs[0].alignment, WD_ALIGN_PARAGRAPH.LEFT)
        with ZipFile(BytesIO(document_bytes)) as archive:
            xml = archive.read('word/document.xml').decode()
            self.assertIn('w:fill="BFBFBF"', xml)
            self.assertIn('w:tblHeader', xml)
            self.assertIn('w:cantSplit', xml)
        book = load_workbook(BytesIO(build_document('xlsx', '2026-10-16', day, classes, settings)))
        self.assertEqual(book.sheetnames, ['1 смена', '2 смена'])
        cells = [c for ws in book for row in ws for c in row if c.value]
        for cls in classes:
            header = next(c for c in cells if c.value == cls['name'])
            self.assertEqual(header.fill.fgColor.rgb, '00BFBFBF')
            self.assertEqual(header.border.bottom.style, 'thin')
        lesson = next(c for c in cells if 'Алгебра ' in str(c.value))
        self.assertEqual(lesson.alignment.horizontal, 'left')
        self.assertEqual(lesson.data_type, 's')
        self.assertTrue(all(ws.oddFooter.right.text is None for ws in book))
        # Short blocks have wider cells and a left inset instead of stretching the whole sheet.
        senior = next(c for c in cells if c.value == '10А')
        first = next(c for c in cells if c.value == '5А')
        self.assertGreater(senior.column, first.column)

    def test_auto_grades_and_manual_groups_preserve_ids_and_class_order(self):
        classes = [{'id': i + 1, 'name': name, 'shift': 1} for i, name in enumerate(['5А', '5Б', '8А', '8Б', '10А', '11А'])]
        blocks = table_blocks(classes, REFERENCE_TEMPLATE)
        self.assertEqual([[c['id'] for c in block] for block in blocks], [[1, 2], [3, 4], [5], [6]])
        classes[-2]['print_group'] = classes[-1]['print_group'] = 'Старшие'
        self.assertEqual([c['id'] for c in table_blocks(classes, REFERENCE_TEMPLATE)[-1]], [5, 6])
        self.assertEqual(len(table_blocks(classes, {**REFERENCE_TEMPLATE, 'group_by_grade': False})), 2)

    def test_long_content_paginates_without_clipping_and_keeps_announcements(self):
        classes, day, settings = reference_fixture()
        settings['template'].update(font_size=11, row_height=20)
        for lesson in day['lessons']:
            lesson['groups'][0]['subject'] = 'Дополнительное занятие по английскому языку и литературе'
            lesson['groups'].append({'subject': 'Иностранный язык — вторая подгруппа', 'teacher': '', 'rooms': ['99'], 'label': 'Группа 2'})
        pages = make_layout('2026-10-16', day, classes, settings, False, [])
        self.assertGreater(len(pages), 2)
        self.assertTrue(all(page['remaining'] >= 0 for page in pages))
        actual_notes = [item['text'] for page in pages for item in page['items'] if item['type'] == 'text']
        for note in day['print_notes']:
            self.assertEqual(actual_notes.count(note['text']), 1)
        actual_headers = [cell for page in pages for item in page['items'] if item['type'] == 'table' for cell in item['rows'][0][2:]]
        self.assertTrue(all(c['name'] in actual_headers for c in classes))

    def test_api_preserves_grouping_and_daily_announcements_in_date_and_weekly_copies(self):
        with tempfile.TemporaryDirectory() as temporary:
            client = create_app(Path(temporary) / 'school.sqlite3').test_client()
            initial = client.get('/api/state').json
            self.assertTrue(initial['capabilities']['reference_print_layout'])
            self.assertIn(REFERENCE_TEMPLATE['name'], [t['name'] for t in initial['templates']])
            classes = initial['classes']
            classes[0]['print_group'] = classes[1]['print_group'] = 'Первая таблица'
            self.assertEqual(client.put('/api/classes', json={'classes': classes}).status_code, 200)
            self.assertEqual(client.get('/api/state').json['classes'][0]['print_group'], 'Первая таблица')
            # Old clients may omit the new field; an ordinary class save must retain it.
            for cls in classes:
                cls.pop('print_group')
            self.assertEqual(client.put('/api/classes', json={'classes': classes}).status_code, 200)
            self.assertEqual(client.get('/api/state').json['classes'][0]['print_group'], 'Первая таблица')
            day = empty_day()
            day['print_notes'] = [{'class_id': 1, 'text': 'Объявление перед первой таблицей'}]
            self.assertEqual(client.put('/api/day/2026-10-15', json={'day': day}).status_code, 200)
            self.assertEqual(client.post('/api/day/2026-10-15/weekly', json={}).status_code, 200)
            for target, payload in [('2026-10-16', {'source': 'saved', 'source_date': '2026-10-15'}), ('2026-10-17', {'source': 'weekly', 'source_weekday': 3})]:
                self.assertEqual(client.post(f'/api/day/{target}/copy', json=payload).status_code, 200)
                self.assertEqual(client.get('/api/day/' + target).json['day']['print_notes'], day['print_notes'])
            bad = deepcopy(day)
            bad['print_notes'].append(day['print_notes'][0])
            self.assertEqual(client.put('/api/day/2026-10-15', json={'day': bad}).status_code, 400)
            self.assertEqual(client.get('/api/day/2026-10-15').json['day']['print_notes'], day['print_notes'])
            settings = initial['settings']
            settings['template'] = REFERENCE_TEMPLATE
            self.assertEqual(client.put('/api/settings', json=settings).status_code, 200)
            self.assertEqual(client.get('/api/state').json['settings']['template'], REFERENCE_TEMPLATE)


if __name__ == '__main__':
    unittest.main()
