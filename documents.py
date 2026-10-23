"""A shared, measured page model for text PDF and editable Office documents."""
from __future__ import annotations

from datetime import date
from io import BytesIO
import math
import os
import re
from pathlib import Path
from xml.sax.saxutils import escape
from scheduling import day_bells

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_LEFT, TA_RIGHT
from reportlab.lib.pagesizes import A3, A4, landscape, portrait
from reportlab.lib.styles import ParagraphStyle
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen.canvas import Canvas
from reportlab.platypus import Paragraph, Table, TableStyle


class LayoutError(ValueError):
    pass


def register_fonts():
    if 'School' in pdfmetrics.getRegisteredFontNames():
        return
    candidates = [
        (os.environ.get('TIMETABLE_FONT', ''), os.environ.get('TIMETABLE_FONT_BOLD', '')),
        ('C:/Windows/Fonts/times.ttf', 'C:/Windows/Fonts/timesbd.ttf'),
        ('/usr/share/fonts/truetype/dejavu/DejaVuSerif.ttf', '/usr/share/fonts/truetype/dejavu/DejaVuSerif-Bold.ttf'),
        ('/usr/share/fonts/truetype/liberation2/LiberationSerif-Regular.ttf', '/usr/share/fonts/truetype/liberation2/LiberationSerif-Bold.ttf'),
    ]
    for regular, bold in candidates:
        if regular and Path(regular).is_file():
            pdfmetrics.registerFont(TTFont('School', regular))
            pdfmetrics.registerFont(TTFont('SchoolBold', bold if bold and Path(bold).is_file() else regular))
            return
    raise LayoutError('Не найден шрифт с кириллицей. Укажите путь к TTF в TIMETABLE_FONT (см. README).')


def text_paragraph(text, font, bold=False, align=TA_CENTER, leading=None):
    style = ParagraphStyle('cell', fontName='SchoolBold' if bold else 'School', fontSize=font,
                           leading=leading or font * 1.17, alignment=align, splitLongWords=True,
                           spaceBefore=0, spaceAfter=0)
    return Paragraph(escape(str(text)).replace('\n', '<br/>') or ' ', style)


def cell_text(lesson, template):
    if not lesson:
        return ''
    lines = []
    if lesson.get('cancelled'):
        lines.append('Урок отменён')
    else:
        for i, group in enumerate(lesson['groups']):
            prefix = group.get('label') or (f'Группа {i + 1}' if len(lesson['groups']) > 1 else '')
            subject = group['subject'] or '(предмет не указан)'
            if prefix:
                subject = f'{prefix}: {subject}'
            if template['show_rooms'] and group['rooms']:
                subject += ' ' + ','.join(group['rooms'])
            lines.append(subject)
            if template['show_teachers'] and group['teacher']:
                lines.append(group['teacher'])
    if lesson.get('note'):
        lines.append(lesson['note'])
    return '\n'.join(lines)


def table_cell_paragraph(value, item, row, column):
    compact = item.get('compact', False)
    return text_paragraph(value, item['font'], bold=row == 0,
                          align=TA_LEFT if compact and row > 0 and column >= 2 else TA_CENTER,
                          leading=item['font'] * 1.08 if compact else None)


def table_blocks(classes, template):
    """Manual table labels take priority; numbered class names can group by grade."""
    groups = {}
    for cls in classes:
        custom = cls.get('print_group', '').strip()
        grade = re.match(r'^\s*(\d{1,2})\s*[А-Яа-яA-Za-zЁё]', cls['name'])
        key = ('manual', custom.casefold()) if custom else ('grade', grade[1]) if grade and template.get('group_by_grade') else ('ordered', '')
        groups.setdefault(key, []).append(cls)
    limit = template['classes_per_table']
    return [values[index:index + limit] for values in groups.values() for index in range(0, len(values), limit)]


def table_flowable(item, template):
    font = item['font']
    padding = item.get('padding', 4)
    data = [[table_cell_paragraph(v, item, i, j) for j, v in enumerate(row)] for i, row in enumerate(item['rows'])]
    table = Table(data, colWidths=item['widths'], rowHeights=item['heights'], hAlign='CENTER')
    table.setStyle(TableStyle([
        ('GRID', (0, 0), (-1, -1), .5, colors.black),
        ('VALIGN', (0, 0), (-1, -1), 'TOP' if item.get('compact') else 'MIDDLE'),
        ('LEFTPADDING', (0, 0), (-1, -1), padding), ('RIGHTPADDING', (0, 0), (-1, -1), padding),
        ('TOPPADDING', (0, 0), (-1, -1), padding), ('BOTTOMPADDING', (0, 0), (-1, -1), padding),
    ]))
    if item.get('compact'):
        table.setStyle(TableStyle([('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#bfbfbf')),
                                   ('VALIGN', (0, 0), (-1, 0), 'MIDDLE')]))
    return table


def make_layout(day_date, day, classes, settings, draft, issues):
    register_fonts()
    template = settings['template']
    compact = template.get('layout') == 'reference'
    size = A3 if template['paper'] == 'A3' else A4
    width, height = (landscape if template['orientation'] == 'landscape' else portrait)(size)
    margin = 28.35
    content_width = width - margin * 2
    font = template['font_size']
    formatted = date.fromisoformat(day_date).strftime('%d.%m.%Y')
    is_error = any(i['severity'] == 'error' for i in issues)
    mark = 'ЧЕРНОВИК — ЕСТЬ ОШИБКИ' if draft and is_error else ('ЧЕРНОВИК' if draft else '')
    lookup = {(l['class_id'], l['slot']): l for l in day['lessons']}
    pages = []

    def text_item(text, size=font, bold=False, align=TA_LEFT, gap=6, text_width=None):
        p = text_paragraph(text, size, bold, align)
        text_width = text_width or content_width
        _, h = p.wrap(text_width, 100000)
        return {'type': 'text', 'text': text, 'font': size, 'bold': bold, 'align': align,
                'height': h, 'gap': gap, 'width': text_width}

    for shift in (1, 2):
        selected = [c for c in classes if c['shift'] == shift]
        if not selected:
            continue
        header = []
        if template['school_name'] and (not compact or not template['show_approval']):
            header.append(text_item(template['school_name'], font + 1, True, TA_CENTER))
        if template['show_approval'] and (not template.get('approval_first_shift_only') or shift == 1):
            approval = template['approval_text']
            if compact and template['school_name']:
                approval += '\n' + template['school_name']
            if template['director']:
                approval += '\n' + ('____________ ' if compact else '') + template['director']
            header.append(text_item(approval, font, False, TA_RIGHT, 5 if compact else 8,
                                    min(245, content_width) if compact else None))
        header.append(text_item(f'Расписание уроков на {formatted} {shift} смена', font + (1 if compact else 3), True, TA_CENTER, 6 if compact else 10))
        if mark:
            header.append(text_item(mark, font + 1, True, TA_CENTER, 8))
        header_height = sum(i['height'] + i['gap'] for i in header)
        capacity = height - 2 * margin - (18 if template.get('show_page_numbers', True) else 0) - header_height
        if capacity < 100:
            raise LayoutError('Шапка занимает страницу. Сократите название школы или строку утверждения.')

        def new_page():
            page = {'shift': shift, 'header': header, 'items': [], 'remaining': capacity,
                    'width': width, 'height': height, 'margin': margin}
            pages.append(page)
            return page

        page = new_page()
        # Notes are paginated as paragraphs, never silently clipped.
        def add_note(text, bold=False, align=TA_LEFT, size=font, gap=5):
            nonlocal page
            if not text:
                return
            for line in text.split('\n'):
                item = text_item(line or ' ', size, bold, align, gap)
                if item['height'] + item['gap'] > capacity:
                    # Divide long prose into words so even a long note fits on pages.
                    words, chunk = line.split(), ''
                    for word in words:
                        candidate = (chunk + ' ' + word).strip()
                        if text_item(candidate, size, bold, align, gap)['height'] > capacity - 10 and chunk:
                            add_note(chunk, bold, align, size, gap)
                            chunk = word
                        else:
                            chunk = candidate
                    if chunk:
                        if text_item(chunk, size, bold, align, gap)['height'] > capacity - 10:
                            raise LayoutError('Слишком длинное слово в примечании. Добавьте пробелы или уменьшите шрифт.')
                        add_note(chunk, bold, align, size, gap)
                    continue
                if item['height'] + item['gap'] > page['remaining']:
                    page = new_page()
                page['items'].append(item)
                page['remaining'] -= item['height'] + item['gap']

        add_note(day['notes'].get(str(shift), ''), compact, TA_CENTER if compact else TA_LEFT,
                 font + 1 if compact else font, 3 if compact else 5)
        bells = day_bells(day, settings['bells'], shift)
        saved_slots = [l['slot'] for l in day['lessons'] if l['class_id'] in {c['id'] for c in selected}]
        min_slot = min([min(bells)] + saved_slots)
        max_slot = max([max(bells)] + saved_slots)
        blocks = table_blocks(selected, template)
        block_index = 0
        while block_index < len(blocks):
            block = blocks[block_index]
            fixed_widths = [16, 58] if compact else [24, 70]
            widths = fixed_widths + [min(template['column_width'], (content_width - sum(fixed_widths)) / len(block))] * len(block)
            rows = [(['', ''] if compact else ['№', 'Время']) + [c['name'] for c in block]]
            for slot in range(min_slot, max_slot + 1):
                rows.append([str(slot), bells.get(slot, 'Время не задано')] +
                            [cell_text(lookup.get((c['id'], slot)), template) for c in block])
            heights = []
            padding = 1.5 if compact else 4
            metrics = {'font': font, 'compact': compact}
            for i, row in enumerate(rows):
                measured = [table_cell_paragraph(v, metrics, i, j).wrap(w - 2 * padding, 100000)[1] + 2 * padding for j, (v, w) in enumerate(zip(row, widths))]
                heights.append(max(template['row_height'] if i else (14 if compact else 23), max(measured)))
            if heights[0] + max(heights[1:]) + 10 > capacity:
                if len(block) > 1:
                    midpoint = math.ceil(len(block) / 2)
                    blocks[block_index:block_index + 1] = [block[:midpoint], block[midpoint:]]
                    continue
                raise LayoutError('Содержимое одного урока превышает высоту страницы. Увеличьте формат бумаги или уменьшите шрифт/текст.')
            block_ids = {c['id'] for c in block}
            announcements = [note['text'] for note in day.get('print_notes', []) if note['class_id'] in block_ids and note['text'].strip()]
            announcement_height = sum(sum(text_item(line or ' ', font + 1, True, TA_CENTER, 3)['height'] + 3 for line in text.split('\n')) for text in announcements)
            complete_height = sum(heights) + 10 + announcement_height
            if complete_height <= capacity and complete_height > page['remaining']:
                page = new_page()
            elif announcement_height + heights[0] + heights[1] + 10 > page['remaining'] and page['items']:
                page = new_page()
            for announcement in announcements:
                add_note(announcement, True, TA_CENTER, font + 1, 3)
            cursor = 1
            while cursor < len(rows):
                if heights[0] + heights[cursor] + 10 > page['remaining']:
                    page = new_page()
                part_rows, part_heights = [rows[0]], [heights[0]]
                remaining = page['remaining'] - heights[0] - 10
                while cursor < len(rows) and heights[cursor] <= remaining:
                    part_rows.append(rows[cursor])
                    part_heights.append(heights[cursor])
                    remaining -= heights[cursor]
                    cursor += 1
                item = {'type': 'table', 'rows': part_rows, 'heights': part_heights, 'widths': widths,
                        'font': font, 'height': sum(part_heights), 'gap': 10,
                        'compact': compact, 'padding': padding}
                page['items'].append(item)
                page['remaining'] -= item['height'] + item['gap']
            if block_index < len(blocks) - 1:
                add_note(template['between_notes'], compact, TA_CENTER if compact else TA_LEFT)
            block_index += 1
    return pages


def pdf_document(pages, template):
    output = BytesIO()
    canvas = Canvas(output, pagesize=(pages[0]['width'], pages[0]['height']), pageCompression=1)
    canvas.setTitle('Расписание уроков')
    canvas.setAuthor(template['school_name'] or 'Школьное расписание')
    for number, page in enumerate(pages, 1):
        canvas.setPageSize((page['width'], page['height']))
        y = page['height'] - page['margin']
        for item in page['header'] + page['items']:
            y -= item['height']
            if item['type'] == 'text':
                p = text_paragraph(item['text'], item['font'], item['bold'], item['align'])
                text_width = item.get('width', page['width'] - page['margin'] * 2)
                p.wrap(text_width, item['height'])
                x = page['width'] - page['margin'] - text_width if item['align'] == TA_RIGHT else page['margin']
                p.drawOn(canvas, x, y)
            else:
                table = table_flowable(item, template)
                table.wrap(0, 0)
                table.drawOn(canvas, (page['width'] - sum(item['widths'])) / 2, y)
            y -= item['gap']
        if template.get('show_page_numbers', True):
            canvas.setFont('School', 8)
            canvas.drawRightString(page['width'] - page['margin'], page['margin'] - 5, f'Страница {number} из {len(pages)}')
        canvas.showPage()
    canvas.save()
    return output.getvalue()


def docx_document(pages, template):
    from docx import Document
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt

    doc = Document()
    section = doc.sections[0]
    section.page_width, section.page_height = Pt(pages[0]['width']), Pt(pages[0]['height'])
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Pt(pages[0]['margin'])
    section.header_distance = section.footer_distance = Pt(12)
    normal = doc.styles['Normal']
    normal.font.name, normal.font.size = 'Times New Roman', Pt(template['font_size'])
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    if template.get('show_page_numbers', True):
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        footer.add_run('Страница ')
        field = OxmlElement('w:fldSimple')
        field.set(qn('w:instr'), 'PAGE')
        footer._p.append(field)
    for page_index, page in enumerate(pages):
        page_items = page['header'] + page['items']
        for item_index, item in enumerate(page_items):
            if item['type'] == 'text':
                p = doc.add_paragraph()
                if page_index and item_index == 0:
                    p.paragraph_format.page_break_before = True
                p.alignment = {TA_LEFT: WD_ALIGN_PARAGRAPH.LEFT, TA_CENTER: WD_ALIGN_PARAGRAPH.CENTER,
                               TA_RIGHT: WD_ALIGN_PARAGRAPH.RIGHT}[item['align']]
                p.paragraph_format.space_after = Pt(item['gap'])
                p.paragraph_format.line_spacing = 1.0
                if item['align'] == TA_RIGHT:
                    p.paragraph_format.left_indent = Pt(page['width'] - 2 * page['margin'] - item.get('width', page['width'] - 2 * page['margin']))
                r = p.add_run(item['text'])
                r.font.size, r.bold = Pt(item['font']), item['bold']
                continue
            table = doc.add_table(rows=0, cols=len(item['widths']))
            table.style = 'Table Grid'
            table.alignment, table.autofit = WD_TABLE_ALIGNMENT.CENTER, False
            borders = OxmlElement('w:tblBorders')
            for side in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
                border = OxmlElement('w:' + side)
                for key, value in [('val', 'single'), ('sz', '4'), ('color', '000000')]:
                    border.set(qn('w:' + key), value)
                borders.append(border)
            table._tbl.tblPr.append(borders)
            for column, width in zip(table.columns, item['widths']):
                column.width = Pt(width)
            for row_index, values in enumerate(item['rows']):
                row = table.add_row()
                # atLeast permits Word to grow rows for its font metrics; no row is split.
                row.height = Pt(item['heights'][row_index])
                properties = row._tr.get_or_add_trPr()
                properties.append(OxmlElement('w:cantSplit'))
                if row_index == 0:
                    properties.append(OxmlElement('w:tblHeader'))
                for column_index, (cell, value, cell_width) in enumerate(zip(row.cells, values, item['widths'])):
                    cell.width = Pt(cell_width)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP if item.get('compact') and row_index > 0 else WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    if item.get('compact') and row_index == 0:
                        shade = OxmlElement('w:shd')
                        shade.set(qn('w:fill'), 'BFBFBF')
                        cell._tc.get_or_add_tcPr().append(shade)
                    margins = OxmlElement('w:tcMar')
                    for side in ('top', 'left', 'bottom', 'right'):
                        edge = OxmlElement('w:' + side)
                        # LibreOffice adds vertical cell margins to trHeight;
                        # reserve padding in the measured row height itself.
                        edge.set(qn('w:w'), '0' if side in ('top', 'bottom') else '30' if item.get('compact') else '60')
                        edge.set(qn('w:type'), 'dxa')
                        margins.append(edge)
                    cell._tc.get_or_add_tcPr().append(margins)
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT if item.get('compact') and row_index > 0 and column_index >= 2 else WD_ALIGN_PARAGRAPH.CENTER
                    p.paragraph_format.line_spacing = 1.08 if item.get('compact') else 1.0
                    r = p.add_run(value)
                    r.font.size, r.bold = Pt(item['font']), row_index == 0
            if item_index < len(page_items) - 1:
                spacer = doc.add_paragraph()
                spacer.paragraph_format.space_after = Pt(0)
                spacer.paragraph_format.line_spacing = Pt(item['gap'])
                spacer.add_run(' ').font.size = Pt(1)
    output = BytesIO()
    doc.save(output)
    return output.getvalue()


def xlsx_document(pages, template):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Side, Font, PatternFill
    from openpyxl.worksheet.pagebreak import Break
    from openpyxl.utils import get_column_letter

    book = Workbook()
    book.remove(book.active)
    line = Side(style='thin', color='000000')
    compact = template.get('layout') == 'reference'
    for shift in (1, 2):
        shift_pages = [p for p in pages if p['shift'] == shift]
        if not shift_pages:
            continue
        ws = book.create_sheet(f'{shift} смена')
        columns = max([len(i['rows'][0]) for p in shift_pages for i in p['items'] if i['type'] == 'table'] or [3])
        widths = max([i['widths'] for p in shift_pages for i in p['items'] if i['type'] == 'table'], key=len, default=[24, 70, 115])
        # A worksheet has global column widths. Use the widest measured width so
        # shorter final blocks never wrap more than their computed row heights.
        class_width = max((w for p in shift_pages for i in p['items'] if i['type'] == 'table' for w in i['widths'][2:]), default=115)
        widths = widths[:2] + [class_width] * (columns - 2)
        content_width = shift_pages[0]['width'] - 2 * shift_pages[0]['margin']
        if compact:
            # A fine grid allows independently sized, centred tables on one sheet.
            # Each visible timetable cell remains an editable merged Excel cell.
            columns = 200
            step = content_width / columns
            # Column widths describe the complete physical grid pitch. Subtracting
            # cell padding for every tiny column would shrink the whole timetable.
            excel_width = step * 96 / 72 / 7
            for c in range(1, columns + 1):
                ws.column_dimensions[get_column_letter(c)].width = excel_width
        else:
            for c, width in enumerate(widths, 1):
                ws.column_dimensions[get_column_letter(c)].width = max(3, (width - 5) / 7)
        ws.sheet_view.showGridLines = False
        ws.page_setup.orientation = template['orientation']
        ws.page_setup.paperSize = ws.PAPERSIZE_A3 if template['paper'] == 'A3' else ws.PAPERSIZE_A4
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
        ws.page_margins.left = ws.page_margins.right = .39
        ws.page_margins.top = ws.page_margins.bottom = .39
        ws.print_options.horizontalCentered = True
        if template.get('show_page_numbers', True):
            ws.oddFooter.right.text = 'Страница &P из &N'
        row = 1
        for p_index, page in enumerate(shift_pages):
            if p_index:
                ws.row_breaks.append(Break(id=row - 1))
            for item in page['header'] + page['items']:
                if item['type'] == 'text':
                    start_col = max(1, columns - round(item.get('width', content_width) / step) + 1) if compact and item['align'] == TA_RIGHT else 1
                    ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=columns)
                    cell = ws.cell(row, start_col, item['text'])
                    cell.data_type = 's'
                    cell.font = Font(name='Times New Roman', size=item['font'], bold=item['bold'])
                    cell.alignment = Alignment(horizontal={TA_LEFT: 'left', TA_CENTER: 'center', TA_RIGHT: 'right'}[item['align']], vertical='center', wrap_text=True)
                    ws.row_dimensions[row].height = item['height'] + item['gap']
                    row += 1
                    continue
                ranges = []
                if compact:
                    position = (content_width - sum(item['widths'])) / 2
                    start_col = round(position / step) + 1
                    for width in item['widths']:
                        position += width
                        end_col = max(start_col, min(columns, round(position / step)))
                        ranges.append((start_col, end_col))
                        start_col = end_col + 1
                else:
                    ranges = [(i, i) for i in range(1, len(item['widths']) + 1)]
                for index, values in enumerate(item['rows']):
                    ws.row_dimensions[row].height = item['heights'][index]
                    for logical_col, (value, (start_col, end_col)) in enumerate(zip(values, ranges)):
                        cell = ws.cell(row, start_col, value)
                        # All inputs are plain text, including values beginning with '='.
                        cell.data_type = 's'
                        cell.font = Font(name='Times New Roman', size=item['font'], bold=index == 0)
                        cell.border = Border(left=line, right=line, top=line, bottom=line)
                        cell.alignment = Alignment(horizontal='left' if compact and index > 0 and logical_col >= 2 else 'center', vertical='top' if compact and index > 0 else 'center', wrap_text=True)
                        if compact and index == 0:
                            cell.fill = PatternFill('solid', fgColor='BFBFBF')
                        if end_col > start_col:
                            ws.merge_cells(start_row=row, start_column=start_col, end_row=row, end_column=end_col)
                    row += 1
                ws.row_dimensions[row].height = item['gap']
                row += 1
        ws.print_area = f'A1:{get_column_letter(columns)}{row - 1}'
    output = BytesIO()
    book.save(output)
    return output.getvalue()


def build_document(fmt, day_date, day, classes, settings, draft=False, issues=()):
    pages = make_layout(day_date, day, classes, settings, draft, issues)
    return {'pdf': pdf_document, 'xlsx': xlsx_document, 'docx': docx_document}[fmt](pages, settings['template'])
